"""API tests for the RBAC management endpoints (D4.3).

The bootstrap admin token (break-glass) authorizes management before any grant
exists — that is the intended bootstrap path.
"""

from __future__ import annotations

import pytest

from tests.conftest import AUTH


@pytest.mark.asyncio
async def test_management_requires_auth(client):
    ac, _ = client
    assert (await ac.get("/api/v1/rbac/roles")).status_code == 401
    assert (await ac.post("/api/v1/rbac/roles", json={"code": "x", "name": "x"})
            ).status_code == 401


@pytest.mark.asyncio
async def test_reseed_then_list(client):
    ac, _ = client
    r = await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)
    assert r.status_code == 200
    assert "admin" in r.json()["roles"]
    roles = (await ac.get("/api/v1/rbac/roles", headers=AUTH)).json()["items"]
    assert {x["code"] for x in roles} >= {"admin", "member"}
    perms = (await ac.get("/api/v1/rbac/permissions", headers=AUTH)).json()
    assert any(p["code"] == "organization.read" for p in perms)


@pytest.mark.asyncio
async def test_roles_export_csv(client):
    ac, _ = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)
    await ac.post("/api/v1/rbac/roles", headers=AUTH,
                  json={"code": "exprole", "name": "Exportable"})
    r = await ac.get("/api/v1/rbac/roles/export", headers=AUTH)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    assert "attachment; filename=roles.csv" in r.headers["content-disposition"]
    lines = r.text.strip().splitlines()
    assert lines[0].startswith("id,code,name,description")
    assert any("exprole" in ln for ln in lines[1:])
    # Filter (q) carries into the export.
    one = await ac.get("/api/v1/rbac/roles/export?q=exprole", headers=AUTH)
    assert len(one.text.strip().splitlines()) == 2  # header + 1 row
    # Reads require auth.
    assert (await ac.get("/api/v1/rbac/roles/export")).status_code == 401


@pytest.mark.asyncio
async def test_roles_export_xlsx(client):
    ac, _ = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)
    await ac.post("/api/v1/rbac/roles", headers=AUTH,
                  json={"code": "xlsxrole", "name": "Spreadsheet"})
    r = await ac.get("/api/v1/rbac/roles/export?format=xlsx", headers=AUTH)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "filename=roles.xlsx" in r.headers["content-disposition"]
    # Body is a real workbook: openpyxl can open it; header row + the new role.
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws["A1"].value == "id"  # frozen header
    codes = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert "xlsxrole" in codes


@pytest.mark.asyncio
async def test_create_set_grants_delete_role(client):
    ac, _ = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)  # catalog
    created = await ac.post("/api/v1/rbac/roles", headers=AUTH,
                            json={"code": "auditor", "name": "Auditor",
                                  "grants": ["organization.read"]})
    assert created.status_code == 201
    assert created.json()["is_system"] is False  # API never mints system roles
    role_id = created.json()["id"]
    # duplicate code in same (global) scope -> 409
    dup = await ac.post("/api/v1/rbac/roles", headers=AUTH,
                        json={"code": "auditor", "name": "Auditor"})
    assert dup.status_code == 409
    upd = await ac.put(f"/api/v1/rbac/roles/{role_id}/permissions", headers=AUTH,
                       json={"codes": ["organization.read", "location.read"]})
    assert upd.status_code == 200
    assert (await ac.delete(f"/api/v1/rbac/roles/{role_id}", headers=AUTH)
            ).status_code == 200


@pytest.mark.asyncio
async def test_get_role_permissions_roundtrip(client):
    ac, _ = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)  # catalog
    created = await ac.post("/api/v1/rbac/roles", headers=AUTH,
                            json={"code": "viewer", "name": "Viewer"})
    role_id = created.json()["id"]
    # No grants yet.
    g0 = await ac.get(f"/api/v1/rbac/roles/{role_id}/permissions", headers=AUTH)
    assert g0.status_code == 200
    assert g0.json()["role_id"] == role_id and g0.json()["codes"] == []
    assert "etag" in g0.json()
    # Set then read back (sorted).
    await ac.put(f"/api/v1/rbac/roles/{role_id}/permissions", headers=AUTH,
                 json={"codes": ["location.read", "organization.read"]})
    g1 = (await ac.get(f"/api/v1/rbac/roles/{role_id}/permissions", headers=AUTH)).json()
    assert g1["codes"] == ["location.read", "organization.read"]
    # Unknown role -> 404.
    assert (await ac.get("/api/v1/rbac/roles/does-not-exist/permissions", headers=AUTH)
            ).status_code == 404


@pytest.mark.asyncio
async def test_role_permissions_optimistic_concurrency(client):
    ac, _ = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)
    rid = (await ac.post("/api/v1/rbac/roles", headers=AUTH,
                         json={"code": "ccrole", "name": "CC"})).json()["id"]
    g = (await ac.get(f"/api/v1/rbac/roles/{rid}/permissions", headers=AUTH)).json()
    etag = g["etag"]
    assert etag

    # Correct etag -> 200, etag rotates.
    ok = await ac.put(f"/api/v1/rbac/roles/{rid}/permissions",
                      headers={**AUTH, "If-Match": etag},
                      json={"codes": ["organization.read"]})
    assert ok.status_code == 200 and ok.json()["etag"] != etag
    # Stale etag -> 409.
    assert (await ac.put(f"/api/v1/rbac/roles/{rid}/permissions",
                         headers={**AUTH, "If-Match": etag},
                         json={"codes": ["location.read"]})).status_code == 409


@pytest.mark.asyncio
async def test_unknown_grant_rejected(client):
    ac, _ = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)
    r = await ac.post("/api/v1/rbac/roles", headers=AUTH,
                      json={"code": "typo", "name": "Typo",
                            "grants": ["organisation.raed"]})  # typo'd code
    assert r.status_code == 422


@pytest.mark.asyncio
async def test_system_role_is_protected(client):
    ac, _ = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)
    roles = (await ac.get("/api/v1/rbac/roles", headers=AUTH)).json()["items"]
    admin_id = next(r["id"] for r in roles if r["code"] == "admin")
    # cannot delete or re-grant a seeded system role
    assert (await ac.delete(f"/api/v1/rbac/roles/{admin_id}", headers=AUTH)
            ).status_code == 409
    assert (await ac.put(f"/api/v1/rbac/roles/{admin_id}/permissions", headers=AUTH,
                         json={"codes": ["organization.read"]})).status_code == 409


@pytest.mark.asyncio
async def test_assignment_records_who_assigned(client):
    ac, db = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)
    roles = (await ac.get("/api/v1/rbac/roles", headers=AUTH)).json()["items"]
    member = next(r["id"] for r in roles if r["code"] == "member")
    acc = (await ac.post("/api/v1/auth/register",
                         json={"password": "Sup3rStr0ng!pw", "email": "w@x.com"})).json()
    assigned = (await ac.post(f"/api/v1/rbac/accounts/{acc['id']}/roles", headers=AUTH,
                              json={"role_id": member, "organization_id": "org-1"})).json()
    # created_by is audited (break-glass principal = bootstrap-admin)
    from app.rbac.models import AccountRole
    async with db.session_factory() as s:
        row = await s.get(AccountRole, assigned["id"])
        assert row.created_by == "bootstrap-admin"


@pytest.mark.asyncio
async def test_assign_and_revoke(client):
    ac, _ = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)
    roles = (await ac.get("/api/v1/rbac/roles", headers=AUTH)).json()["items"]
    admin_id = next(r["id"] for r in roles if r["code"] == "admin")
    acc = (await ac.post("/api/v1/auth/register",
                         json={"password": "Sup3rStr0ng!pw", "email": "u@x.com"})).json()
    assigned = await ac.post(f"/api/v1/rbac/accounts/{acc['id']}/roles", headers=AUTH,
                             json={"role_id": admin_id, "organization_id": "org-1"})
    assert assigned.status_code == 201
    assignment_id = assigned.json()["id"]
    listed = (await ac.get(f"/api/v1/rbac/accounts/{acc['id']}/roles",
                           headers=AUTH)).json()
    assert len(listed) == 1 and listed[0]["organization_id"] == "org-1"
    assert (await ac.delete(
        f"/api/v1/rbac/accounts/{acc['id']}/roles/{assignment_id}", headers=AUTH)
    ).status_code == 200


@pytest.mark.asyncio
async def test_assign_with_unit_site_scope(client):
    """The assign endpoint accepts + persists the full org/unit/site scope tuple
    — the exact contract the ScopePicker UI relies on (F.5 parity)."""
    ac, _ = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)
    roles = (await ac.get("/api/v1/rbac/roles", headers=AUTH)).json()["items"]
    member = next(r["id"] for r in roles if r["code"] == "member")
    acc = (await ac.post("/api/v1/auth/register",
                         json={"password": "Sup3rStr0ng!pw", "email": "scope@x.com"})).json()
    r = await ac.post(f"/api/v1/rbac/accounts/{acc['id']}/roles", headers=AUTH,
                      json={"role_id": member, "organization_id": "org-1",
                            "org_unit_id": "unit-1", "site_id": "site-1"})
    assert r.status_code == 201, r.text
    listed = (await ac.get(f"/api/v1/rbac/accounts/{acc['id']}/roles", headers=AUTH)).json()
    assert listed[0]["organization_id"] == "org-1"
    assert listed[0]["org_unit_id"] == "unit-1"
    assert listed[0]["site_id"] == "site-1"


@pytest.mark.asyncio
async def test_bulk_assign_role(client):
    ac, _ = client
    await ac.post("/api/v1/rbac/admin/reseed?profile=empty", headers=AUTH)
    roles = (await ac.get("/api/v1/rbac/roles", headers=AUTH)).json()["items"]
    member = next(r["id"] for r in roles if r["code"] == "member")
    a1 = (await ac.post("/api/v1/auth/register",
                        json={"password": "Sup3rStr0ng!pw", "email": "b1@x.com"})).json()
    a2 = (await ac.post("/api/v1/auth/register",
                        json={"password": "Sup3rStr0ng!pw", "email": "b2@x.com"})).json()

    r = await ac.post("/api/v1/rbac/accounts/bulk-roles", headers=AUTH,
                      json={"account_ids": [a1["id"], a2["id"], a1["id"]],  # dupe ignored
                            "role_id": member, "organization_id": "org-1"})
    assert r.status_code == 201
    body = r.json()
    assert set(body["assigned"]) == {a1["id"], a2["id"]}
    assert body["errors"] == []
    # Both now carry the role.
    for aid in (a1["id"], a2["id"]):
        listed = (await ac.get(f"/api/v1/rbac/accounts/{aid}/roles", headers=AUTH)).json()
        assert any(x["role_id"] == member for x in listed)

    # Unknown role -> whole call fails (404); empty -> 422.
    assert (await ac.post("/api/v1/rbac/accounts/bulk-roles", headers=AUTH,
                          json={"account_ids": [a1["id"]], "role_id": "nope"})
            ).status_code == 404
    assert (await ac.post("/api/v1/rbac/accounts/bulk-roles", headers=AUTH,
                          json={"account_ids": [], "role_id": member})
            ).status_code == 422


@pytest.mark.asyncio
async def test_assign_unknown_role_404(client):
    ac, _ = client
    acc = (await ac.post("/api/v1/auth/register",
                         json={"password": "Sup3rStr0ng!pw", "email": "v@x.com"})).json()
    r = await ac.post(f"/api/v1/rbac/accounts/{acc['id']}/roles", headers=AUTH,
                      json={"role_id": "does-not-exist"})
    assert r.status_code == 404
